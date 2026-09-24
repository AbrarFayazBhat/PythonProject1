##MY DATA MY STORY 
#importing openpyxl and math library
import openpyxl
import math
wb=openpyxl.load_workbook("12607000.xlsx", read_only=True,data_only=True)
ws=wb["Daily Log"]
#CALCULATE TECH PRODUCTIVITY INDEX
def calculate_TPI(column,valid_days,end_row):
    values=[]
    for row in range(6,end_row+1):
        value=ws.cell(row,column).value
        if value is not None:
            values.append(value)
    return round(math.fsum(values)/valid_days,2)
##CALUCULATE ACADEMIC ACTIVITY INDEX
def calculate_AAI(study_column,class_column,valid_days,end_row):
    study_minutes=[]
    for row in range(6,end_row+1):
        value=ws.cell(row,study_column).value
        if value is not None:
            study_minutes.append(value)
    class_minutes=[]
    for row in range(6,end_row+1):
        value=ws.cell(row,class_column).value
        if value is not None:
            class_minutes.append(value)
    return round(math.fsum(study_minutes+class_minutes)/valid_days,2)
##CALCULATE PHYSICAL ACTIVITY INDEX
def calculate_PHAI(column,valid_days,end_row):
    fitness_minutes=[]
    for row in range(6,end_row+1):
        value=ws.cell(row,column).value
        if value is not None:
            fitness_minutes.append(value)
    return round(math.fsum(fitness_minutes)/valid_days,2)
##CALCULATE SLEEP AND RECOVERY INDEX
def calculate_SRI(column,valid_days,end_row):
    sleep_minutes=[]
    for row in range(6,end_row+1):
        value=ws.cell(row,column).value
        if value is not None:
            sleep_minutes.append(value)
    return round(math.fsum(sleep_minutes)/valid_days,2)
##CALCULATE ACTIVITY BALANCE INDEX
def calculate_ABI(freetime,valid_days,end_row):
    free_minutes=[]
    for row in range(6,end_row+1):
        value=ws.cell(row,freetime).value
        if value is not None:
            free_minutes.append(value)
    return round(math.fsum(free_minutes)/valid_days,2)

##CALCULATE TIME UTILIZATION INDEX
def calculate_TUI(column,valid_days,end_row):
    utilization_minutes=[]
    for row in range(6,end_row+1):
        value=ws.cell(row,column).value 
        if value is not None:
            utilization_minutes.append(value)
    return round(math.fsum(utilization_minutes)/valid_days,2)
##CALCULATE EXPERIENCE INDEX(EI)
def calculate_EI(feeling,satisfication,energy,valid_days,end_row):
    feeling_level=[]
    # CALCULATE SUM OF FEELING LEVELS
    for row in range(6,end_row+1):
        values=ws.cell(row,feeling).value
        if values is not None:
            if values=="Excellent":
                feeling_level.append(5)
            elif values=="Good":
               feeling_level.append(4)
            elif values=="Neutral":
               feeling_level.append(3)
            elif values=="Low":
              feeling_level.append(2)
            else:
               feeling_level.append(1)
    satisfication_level=[]
    # CALCULATING SUM OF SATISFICATION LEVELS
    for row in range(6,end_row+1):
       values=ws.cell(row,satisfication).value
       if values is not None:
            if values=="Very Satisfied":
                 satisfication_level.append(5)
            elif values=="Satisfied":
                satisfication_level.append(4)
            elif values=="Neutral":
                satisfication_level.append(3)
            elif values=="Unsatisfied":
                satisfication_level.append(2)
            else:
                satisfication_level.append(1)
    energy_level=[]
    # CALCULATING SUM OF ENERGY LEVELS
    for row in range(6,end_row+1):
        values=ws.cell(row,energy).value
        if values is not None:
            if values=="High":
                energy_level.append(3)
            elif values=="Medium":
                energy_level.append(2)
            else:
                energy_level.append(1)
    return round(math.fsum(feeling_level+satisfication_level+energy_level)/(3*valid_days),2)
##DATA CONTINUITY INDEX
def calculate_DCI(valid_days,expected_days):
    return round((valid_days/expected_days)*100,2)

##calling all functions
tpi=calculate_TPI(5,36,42)
print(f'Tech Productivity Index is: {tpi}')

aai=calculate_AAI(4,6,36,42)
print(f'Academic Activity Index is: {aai}')

Phai=calculate_PHAI(3,36,42)
print(f'Physical Activity Index is: {Phai}')

Sri=calculate_SRI(2,36,42)
print(f'Sleep and  Recovery Index is: {Sri}')

abi=calculate_ABI(10,36,42)
print(f'Activity Balance Index is: {abi}')

tui=calculate_TUI(9,36,42)
print(f'Time Utilization Index is: {tui}')

ei=calculate_EI(11,12,13,36,42)
print(f'Experience Index is: {ei}')

dci=calculate_DCI(36,36)
print(f'Data Continuity index is: {dci}')

# CALCULATING PERSONAL ACTIVITY INDEX
PAI=round((0.15*tpi)+(0.20*aai)+(0.15*Phai)+(0.15*tui)+(0.2*Sri)+(0.10*ei)+(0.05*dci),2)
print(f'Personal Activity Index is: {PAI}')
